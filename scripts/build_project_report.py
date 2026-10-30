"""Build the academic PDF and editable DOCX from supplied PDFs and recorded evidence.

Uses already-installed pypdfium2, Pillow, reportlab and python-docx. No network,
database access, provider calls or application changes. Run from any directory.
"""
from hashlib import sha256
from pathlib import Path
from xml.sax.saxutils import escape
import json
import math

import pypdfium2 as pdfium
from PIL import Image as PILImage, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, Frame, Image, PageBreak,
                               PageTemplate, Paragraph, Spacer, Table, TableStyle)
from reportlab.platypus.tableofcontents import TableOfContents
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "var/report-build"
OUT = ROOT / "reports"
FONT = Path("C:/Windows/Fonts")
INK, GREEN, LIGHT = "#203b34", "#3b6c56", "#edf3ef"
SOURCE_FILES = ["Assignment_10_Report.pdf", "sampe output.pdf", "plan-proposal.md",
                "TermsLens_Project_Plan.md"]


def font(size, bold=False):
    return ImageFont.truetype(str(FONT / ("arialbd.ttf" if bold else "arial.ttf")), size)


def centered(draw, xy, text, size=28, bold=False, fill=INK):
    draw.multiline_text(xy, text, font=font(size, bold), fill=fill, anchor="mm", align="center", spacing=10)


def arrow(draw, points, label=None, label_xy=None, color=GREEN):
    draw.line(points, fill=color, width=5)
    a, b = points[-2], points[-1]
    angle = math.atan2(b[1] - a[1], b[0] - a[0])
    wings = [(b[0] - 17 * math.cos(angle + offset), b[1] - 17 * math.sin(angle + offset)) for offset in (-.5, .5)]
    draw.polygon([b, *wings], fill=color)
    if label:
        centered(draw, label_xy, label, 22)


def box(draw, bounds, title, body="", fill=LIGHT):
    draw.rounded_rectangle(bounds, radius=20, fill=fill, outline=GREEN, width=3)
    x1, y1, x2, y2 = bounds
    centered(draw, ((x1+x2)/2, y1+45), title, 30, True)
    if body:
        centered(draw, ((x1+x2)/2, (y1+y2)/2+22), body, 24)


def prepare_images():
    for filename, prefix in [("Assignment_10_Report.pdf", "cover-reference"), ("sampe output.pdf", "output")]:
        pdf = pdfium.PdfDocument(ROOT / filename)
        for i in range(1 if prefix == "cover-reference" else len(pdf)):
            page = pdf[i]
            bitmap = page.render(scale=2)
            bitmap.to_pil().convert("RGB").save(WORK / f"{prefix}-{i+1}.png")
            bitmap.close()
            page.close()
        pdf.close()
    cover = PILImage.open(WORK / "cover-reference-1.png").convert("RGB")
    d = ImageDraw.Draw(cover)
    d.rectangle((300, 130, 1040, 218), fill="white")
    centered(d, (612, 170), "MINI PROJECT", 36, True, "black")
    centered(d, (612, 218), "TermsLens — Consumer Policy Analysis", 23, False, "black")
    cover.save(WORK / "cover.png")
    crops = {
        "overview": (1, (82, 302, 1144, 1522)),
        "highlights": (2, (85, 110, 1144, 1520)),
        "evidence": (4, (88, 147, 1144, 1438)),
        "filters": (3, (88, 800, 1144, 1530)),
        "controls": (5, (88, 961, 1144, 1468)),
        "deletion": (6, (82, 57, 1144, 410)),
    }
    for name, (page, bounds) in crops.items():
        with PILImage.open(WORK / f"output-{page}.png") as source:
            source.crop(bounds).save(WORK / f"{name}.png")

    img = PILImage.new("RGB", (1400, 820), "white")
    d = ImageDraw.Draw(img)
    centered(d, (700, 37), "TermsLens: application and trust boundaries", 32, True)
    box(d, (40, 100, 410, 280), "React browser", "Input, consent, history\nEvidence and score views")
    box(d, (515, 100, 885, 280), "FastAPI service", "Authentication, validation\nOwned reads and jobs")
    box(d, (990, 100, 1360, 280), "MySQL", "Sessions, job queue\nCaptures, reports, quotas")
    arrow(d, [(410, 190), (515, 190)], "HTTP", (462, 158))
    arrow(d, [(885, 190), (990, 190)], "SQL", (938, 158))
    box(d, (40, 470, 410, 650), "Private snapshots", "Captured source files\nOutside public assets")
    box(d, (515, 470, 885, 650), "Python worker", "Capture, rules, AI review\nScore and publication")
    box(d, (990, 470, 1360, 650), "Groq boundary", "Selected clauses/context\nValidated response only")
    arrow(d, [(1165, 280), (1165, 365), (700, 365), (700, 470)], "Claim jobs; recover saved work", (916, 334))
    arrow(d, [(885, 510), (935, 510), (935, 310), (990, 310), (990, 260)], "Publish", (995, 413))
    arrow(d, [(515, 558), (410, 558)], "Files", (462, 526))
    arrow(d, [(885, 560), (990, 560)], "Consent", (938, 534))
    arrow(d, [(990, 605), (885, 605)], "Evidence", (938, 635))
    centered(d, (700, 727), "The browser reads saved reports through the API. Groq has no database or filesystem access.", 23)
    centered(d, (700, 773), "One normal worker • owner-scoped records • durable job and quota state", 24, True)
    img.save(WORK / "architecture.png")

    img = PILImage.new("RGB", (1400, 770), "white")
    d = ImageDraw.Draw(img)
    centered(d, (700, 38), "Recovery is based on durable state", 32, True)
    box(d, (45, 145, 345, 315), "Queued", "Committed input\nOwner + submission key")
    box(d, (525, 145, 865, 315), "Running", "Lease + stage\nAt most three attempts")
    box(d, (1060, 70, 1370, 225), "Done / partial", "Atomic report + job link")
    box(d, (1060, 270, 1370, 425), "Failed", "No usable report\nBounded error code", "#fbefea")
    arrow(d, [(345, 230), (525, 230)], "Claim", (435, 198))
    arrow(d, [(865, 190), (1060, 147)], "Publish", (947, 140))
    arrow(d, [(865, 265), (1060, 340)], "Fatal / exhausted", (960, 246))
    arrow(d, [(600, 315), (600, 392), (190, 392), (190, 315)], "Expired lease; attempts remain", (407, 432))
    labels = [("Reserve", "Persist capacity"), ("Dispatch", "Mark sending once"), ("Save response", "Reconcile usage"), ("Validate", "Publish or abstain")]
    for i, (title, body) in enumerate(labels):
        x = 40+i*345
        box(d, (x, 525, x+280, 665), title, body)
        if i < 3: arrow(d, [(x+280, 595), (x+345, 595)])
    centered(d, (700, 723), "Unknown provider outcome: retain the reservation and do not resend automatically.", 24, True)
    img.save(WORK / "recovery.png")


def P(text): return ("p", text)
def H(text): return ("h", text)
def B(text): return ("bullet", text)
def T(headers, rows, widths=None): return ("table", headers, rows, widths)
def F(name, caption, max_height): return ("figure", name, caption, max_height)
def C(text): return ("code", text)


UNITS = [
    ("Abstract", [
        P("TermsLens is a full-stack mini project that helps consumers read terms, privacy notices and related online policies through evidence-backed findings. Instead of presenting a single unexplained score, it separates concerning permissions, favorable protections and unresolved wording, and connects each supported finding to its exact source passage. The central question is practical: what does the policy permit, what protection does it promise, and where is that statement written?"),
        P("The application combines a React and TypeScript interface, a FastAPI service, a MySQL persistence layer and a separate Python worker. Local rules handle explicitly supported wording across eight clause groups. Consented Hybrid review can submit a bounded selection of unresolved clauses to Groq. Scoring, coverage and publication remain controlled by application code; the model does not assign the overall grade."),
        P("This report documents the implemented architecture, security controls, rate limiting, failure handling and user workflow. Images come from the supplied six-page printed output. That observed run requested Hybrid but saved a Rules-only result: 14 of 20 substantive clauses were rated, with nine concerns and five favorable findings. It demonstrates a useful local baseline and visible degradation, while successful live AI acceptance and general-policy quality remain unverified."),
        T(["Project", "Details"], [
            ["Title", "Mini Project — TermsLens"], ["Student", "KAMAL SUNDAR SK [26120137]"],
            ["Faculty guide", "Dr. Rajalakshmi"], ["Course", "Full Stack Programming"],
            ["Institution", "Shiv Nadar University Chennai"], ["Report date", "30 October 2026"],
            ["Demonstrated result", "70% clause coverage; 14 rule findings; six unresolved clauses"],
        ], [130, 374]),
        P("Keywords: consumer policies, explainable analysis, rule-based systems, bounded AI, FastAPI, MySQL, consent, evidence, recovery."),
    ], False),
    ("Contents", [("toc",)], False),
    ("1. Introduction and objectives", [
        H("1.1 Problem statement"),
        P("Online services commonly ask users to accept long policies before creating an account, purchasing a subscription or uploading content. Important permissions can be distributed across sections on advertising, licensing, cancellation and account closure. A favorable sentence can also be misunderstood: a promise not to sell information does not necessarily prohibit sharing it with advertisers, and the ability to cancel a subscription does not itself guarantee a refund."),
        P("TermsLens addresses this reading problem by turning supported policy passages into navigable findings. It assists inspection of the text; it does not determine a company's hidden intentions, establish legal enforceability or replace professional advice. Unfamiliar language and missing context are displayed as unresolved instead of being silently classified as harmless. [1, 2]"),
        H("1.2 Project objectives"),
        B("Provide an authenticated workflow for submitting one English policy and reopening the saved analysis."),
        B("Identify supported consumer risks and protections across eight defined clause groups."),
        B("Explain each finding using an exact quote, a source label and a severity rating."),
        B("Use optional external AI only with consent, verified configuration and a reservable budget."),
        B("Preserve immutable policy versions and reports so later comparisons have a traceable basis."),
        B("Make uncertainty, unavailable AI, failures and incomplete coverage visible to the user."),
        H("1.3 Scope and academic contribution"),
        P("The project is a local educational demo. Its contribution is the integration of interpretation with durable jobs, ownership checks, budget accounting, evidence validation and usable report navigation. The implemented system supports Terms, Privacy, Cookies and Refund policy scopes. The demonstration policy is authored engineering material, not a human-labeled benchmark. General-purpose legal interpretation, broad semantic coverage and production deployment are outside the evidence established by this submission."),
    ], True),
    ("2. Features and consumer workflow", [
        P("The interface starts with authentication and a policy form. The user chooses the policy type, pastes text or enters a public URL, and selects a review mode. Hybrid is the default when available, but its consent checkbox remains unchecked. The API returns a job identifier; the UI follows actual processing stages and opens the saved report when publication finishes."),
        T(["Clause group", "Consumer question"], [
            ["Sharing / selling", "Can information be sold or passed to another organization?"],
            ["Retention / deletion", "How long can information be retained, and can it be deleted?"],
            ["Tracking / consent", "What tracking occurs, and does it require consent?"],
            ["Content licensing / AI training", "What rights are granted over uploaded content?"],
            ["Renewal / refunds / cancellation", "Can charges continue, and can payments be recovered?"],
            ["Unilateral changes", "Can the service change the agreement without notice?"],
            ["Arbitration / class waivers", "How are disputes handled and collective claims restricted?"],
            ["Termination / export", "Can access end abruptly, and can data be taken away?"],
        ], [185, 319]),
        H("Report and lifecycle features"),
        P("The report begins with score bounds, coverage, key concerns, favorable findings and unresolved counts. Users can filter clauses by category, rating or source, inspect exact evidence, switch stored scoring profiles and compare the saved result with its rules-only baseline. Longer lists are paginated rather than displayed as one uninterrupted block."),
        P("Private history supports reopening reports. URL discovery offers up to four candidates for explicit selection. Refresh and reanalysis preserve separate source versions and interpretation revisions. Compatible reports can be compared; deletion uses confirmation and a recoverable cleanup status. These features are implemented, with the remaining runtime acceptance work recorded separately from the browser and unit-test evidence. [2–5]"),
    ], True),
    ("3. System architecture", [
        F("architecture", "Figure 1. Architecture drawn from the implemented API, worker, database and provider boundaries; this is a technical diagram, not a runtime screenshot.", 295),
        P("React 19 and TypeScript provide the browser interface, with Vite used for development and production assets. FastAPI validates requests and applies authentication, ownership, origin and CSRF checks. SQLAlchemy and Alembic provide relational persistence and schema migrations over MySQL 8. The development setup serves the built interface from Uvicorn on port 8000; Vite can run separately on port 5173. [2–4]"),
        P("Analysis does not run inside the submission response. The API commits a job before returning HTTP 202. One normal Python worker claims queued work, captures the source, segments clauses, runs local rules, optionally coordinates AI, calculates the result and publishes the report. The browser polls the owned job rather than assuming progress from elapsed time."),
        P("The worker owns the external-processing boundary. Groq receives only selected clauses and required context. It has no database, filesystem or tool access. Raw fetched snapshots remain in private storage outside public assets. Saved reports are read through owner-checked API routes."),
    ], True),
    ("4. Data design and persistence", [
        P("The relational design separates source material, interpretation and workflow state. A logical document may have several immutable captures, and each capture may have several analysis revisions. This distinction permits the system to report whether wording changed or the same wording was interpreted under different settings. [2, 3]"),
        T(["Record group", "Responsibility"], [
            ["Users and sessions", "Account identity, password hashes and hashed expiring session tokens."],
            ["Documents and versions", "Ownership, source identity, ordered captures, cleaned text, hashes and capture metadata."],
            ["Jobs", "Requested mode, consent, idempotency key, state, stage, attempts, lease and result link."],
            ["Analyses", "Engine revisions, effective mode, scores, coverage, baseline and provenance."],
            ["Clauses and findings", "Exact source offsets, substantive status, uncertainty and rule/AI evidence."],
            ["Category scores", "Persisted category-level interpretation and score inputs."],
            ["AI calls and budget windows", "Durable attempts, payload identity, responses, usage and reservations."],
            ["Comparisons and deletions", "Saved comparisons, tombstones and recoverable cleanup receipts."],
            ["Fetch-host records", "Shared host leases and request spacing for public URL operations."],
        ], [164, 340]),
        H("Consistency boundaries"),
        P("An owner-scoped idempotency key returns the existing job for an identical submission; reusing the key with different input is rejected. Publication commits the report, its job result link and terminal status atomically. A stale worker cannot publish after losing its lease. Comparisons require compatible source types, engine settings and effective modes."),
        P("Evidence offsets refer to the immutable cleaned capture using Unicode code points. The browser converts or slices these carefully so a non-BMP character, such as an emoji before a quote, does not shift the highlight. Profile switching uses stored findings and causes no additional provider call. Real MySQL restart, isolation and concurrency acceptance remain pending even though local persistence checks have passed."),
    ], True),
    ("5. Analysis, evidence and scoring", [
        H("5.1 Rule-first interpretation"),
        P("The current engine segments text into clauses, excludes recognized headings and checks a versioned set of supported statements. These initial engineering rules cover all eight groups, but their wording coverage is narrow. Actors, negation, qualifications and conflicting statements must not be discarded to manufacture a match. Supported findings store a category, severity, explanation and exact quote; unsupported clauses remain unrated."),
        H("5.2 Bounded Hybrid review"),
        P("Unresolved clauses may be selected for Groq only after explicit consent and readiness checks. The payload includes IDs and context and treats the policy as untrusted data. The response must satisfy a strict schema, refer to selected IDs, use valid category/rating combinations and quote the supplied clause exactly. Uncertain, conflicting, duplicate or fabricated evidence is rejected. Source and model provenance distinguish accepted AI findings from local rules."),
        H("5.3 Deterministic score calculation"),
        P("Higher scores are more favorable; higher clause severity indicates greater concern. Severity runs from 0 to 5: favorable 0, neutral 1, concerning 2–3 and invasive 4–5. Each applicable category uses supported severity maxima rather than averaging repeated favorable statements against a serious concern. Unknown evidence widens the possible score interval. [1, 4]"),
        T(["Measure", "Interpretation"], [
            ["Coverage", "Rated substantive clauses / all substantive clauses. This is not accuracy."],
            ["Upper score bound", "100 minus 20 × the weighted mean of known category severity maxima."],
            ["Lower score bound", "Uses worst-case severity for category uncertainty; fixed scope is retained."],
            ["Grade requirement", "At least 80% coverage, consistent scope and both bounds in the same grade band."],
            ["Critical caps", "Accepted data sale caps at 74; perpetual licensing plus mandatory arbitration caps at 59."],
        ], [134, 370]),
        P("Terms uses all eight groups; Privacy uses groups 1, 2, 3, 4 and 6; Cookies uses 1, 2, 3 and 6; Refund uses 5 and 6. Balanced and Privacy-first views apply fixed weights to this scope. The model never assigns the overall grade. A report with no supported findings displays ‘Not enough evidence’ rather than treating an unknown upper bound as a favorable verdict."),
    ], True),
    ("6. Security, privacy and trust boundaries", [
        H("6.1 Authentication and session protection"),
        P("Passwords are hashed using Argon2id; the current hasher uses time cost 2, memory cost 19,456 KiB and parallelism 1. Random session tokens are generated using the platform secrets module, while only their SHA-256 hashes are stored in the database. Sessions expire after 12 hours by default. Logout removes the session record. Login failure messages do not identify which credential was incorrect. [4]"),
        P("The session cookie is HttpOnly, SameSite=Strict and restricted to /api. State-changing authenticated requests require an allowed Origin and a matching CSRF token checked with a constant-time comparison. In the local HTTP demo, Secure cookies are disabled; an HTTPS deployment must enable them. This report does not claim that the local configuration is a production security audit."),
        H("6.2 Authorization and private data"),
        P("Report, job, history, evidence, comparison and document operations derive the owner from the authenticated session. A document identifier supplied by a client is not sufficient authorization. Deletion tombstones block new attachments or refresh work while cleanup is pending. Test CLI settings select the isolated test database and clear Groq credentials, model, readiness path and enable flag."),
        H("6.3 Input, rendering and diagnostics"),
        P("JSON models validate request types and lengths. A body-limit middleware bounds chunked requests before JSON parsing, including requests without a Content-Length header. React renders policy text as escaped content; fetched raw HTML is never rendered as a document page. Responses include nosniff, same-origin referrer policy and frame-denial headers; API responses use no-store caching."),
        P("Secrets remain in private .env or host environment storage. Git ignores private configuration, snapshots, local databases and generated caches. Diagnostics expose fixed codes and correlation identifiers, not passwords, API keys or policy bodies. The report and its build script use only supplied documents and recorded project evidence; they make no provider request."),
    ], True),
    ("6.4 Rate limiting and resource protection", [
        P("Rate limits serve different purposes at different boundaries. The API limiter reduces repeated user/IP requests within one API process. The database ledger protects provider capacity across jobs and restarts. Neither mechanism is a substitute for the other. The values below describe current implementation and the recorded demo account, not universal Groq limits. [3, 4, 6]"),
        T(["Boundary", "Current bound"], [
            ["Registration / login", "5 registration requests/minute/IP; 10 login requests/minute/IP."],
            ["New submissions", "10/minute/IP and 10/minute/user; at most five queued/running jobs per owner."],
            ["Discovery / comparison", "Discovery: 3/minute/IP and user. Comparison: 10/minute/IP and user."],
            ["Request / policy input", "1 MiB request body; policy ≤200,000 characters and ≤25,000 words; ≤2,000 clauses."],
            ["Public URL operation", "8 requests including redirects/robots, 3 redirects, 2 MiB per decompressed response, 60 seconds."],
            ["AI per analysis", "At most 2 requests and 12 selected clauses, at most 6 clauses per batch; 1 retry maximum."],
            ["AI token reservations", "≤6,000 input/request; shared 10,000 input and 2,000 output across attempts. Output ≤1,000/request."],
            ["Recorded Groq account", "openai/gpt-oss-20b: 30 RPM, 1,000 RPD, 8,000 TPM and 200,000 TPD."],
            ["API limiter storage", "Bounded to 10,000 active buckets; cannot evict a live bucket to create fresh capacity."],
        ], [148, 356]),
        H("Public URL and provider controls"),
        P("URL fetching permits public HTTP(S) targets on ports 80/443, validates all DNS answers, pins the selected address while preserving Host/TLS checks, and revalidates redirects. Private, loopback, link-local and other forbidden addresses are rejected to reduce server-side request forgery risk. Robots rules, shared host leases and request spacing limit repeated fetching. Dynamic or blocked pages can be pasted manually."),
        P("The recent per-request increase from 5,000 to 6,000 does not raise the shared analysis total or account quotas. Token reservations use a conservative estimator and must cover context and output. Unknown or expired readiness permits zero calls. A larger first request can leave insufficient capacity for a retry; the ledger must refuse that retry. API request counters are in memory and reset with the API process, so a scaled deployment would need coordinated ingress limits."),
    ], False),
    ("7. Failure engineering and diagnosis", [
        P("Failure engineering makes unsuccessful paths explicit and recoverable. TermsLens distinguishes configuration failure, unavailable external review, incomplete interpretation and an unusable job. A terminal request is not automatically proof that AI ran or that the policy received a reliable overall assessment. [2–5]"),
        T(["Observed or handled failure", "Engineering response and lesson"], [
            ["Missing dependencies / MySQL login 1045", "Safe diagnostics separate installation, login and schema checks. An administrator's successful login does not validate an application account."],
            ["Worker appears to hang", "The worker remains open while waiting for queued jobs. Startup messaging explains that behavior; Ctrl+C stops it cleanly."],
            ["AI disabled in one terminal", "PowerShell process environment overrides .env. Set the enable flag in both API and worker terminals before launching."],
            ["Readiness expired", "Expose a fixed reason. Explicit renewal rechecks model/tokenizer, archives the old record and preserves quota history."],
            ["Context exceeds budget", "Abstain without truncating an exception. Earlier 97- and 290-clause policies selected zero clauses for AI."],
            ["Provider timeout / unknown outcome", "Retain reservations and record uncertainty. Do not automatically resend a possibly completed call."],
            ["Poor report readability", "Move the overview first, group risks/protections/unknowns, paginate clauses and use exact-evidence dialogs."],
        ], [167, 337]),
        H("Diagnose the boundary that failed"),
        P("HTTP 200 from /api/health confirms the health response succeeded; it does not by itself mean AI is available. Inspect external_review.available and its fixed reason. Similarly, a successful model lookup verifies availability of the configured model, not an inference result. Request IDs connect API responses to jobs, while job IDs connect stage events, capture versions and published results."),
        P("The project preserves fallback reasons and requested/effective modes in saved reports. This makes a Rules-only result useful without mislabeling it as Hybrid success. General context selection remains a known blocker: the current rule guard can suppress the entire document and the AI payload can collect unrelated guarded paragraphs."),
    ], True),
    ("7.1 Recovery and child-process handling", [
        F("recovery", "Figure 2. Job recovery and provider attempt ordering. Local tests cover these transitions; complete real-MySQL crash/race acceptance remains pending.", 285),
        P("A worker claims a queued job with a lease and advances its processing stage. Recoverable interruptions may return work to the queue while attempts remain. The initial design permits one active worker and three job attempts. Losing or replacing a lease revokes publication rights so a stale process cannot create a second report."),
        P("Before calling Groq, the coordinator reserves capacity and persists the attempt. Dispatch is marked once; the response and usage are saved before interpretation. Compatible saved responses can be recovered without a fresh call. Unknown outcomes keep their charge, and completed or interrupted calls are not automatically replayed. A known 429 rejection may qualify for one bounded retry, subject to all remaining request/token quotas."),
        P("External work runs in bounded child processes. Parent cleanup kills and reaps children after timeout, communication failure or interruption; the Groq child also has an independent deadline watchdog. Credentials and payloads travel through standard input rather than command-line arguments. URL children use bounded request/result envelopes. Test workers clear provider settings so offline acceptance cannot consume production Groq capacity."),
    ], False),
    ("8. Output and consumer interpretation", [
        P("The supplied sampe output.pdf is a six-page browser print captured on 30 October 2026 at approximately 09:20. Figures 3–8 are cropped from those pages for readability; their displayed results have not been changed. The captured report demonstrates the local baseline and a visible external-review fallback. [7]"),
        T(["Visible field", "Observed value and meaning"], [
            ["Policy type", "Privacy. The broad consumer example also contains terms-of-service categories."],
            ["Requested / saved mode", "Hybrid / Rules-only. External review is shown as unavailable."],
            ["Evidence coverage", "70%: 14 of 20 substantive clauses have supported findings."],
            ["Finding groups", "Nine clauses need attention; five favorable findings; six unresolved clauses."],
            ["Displayed score", "0–52 / 100, provisional, with no single letter grade."],
            ["Scope warning", "Some findings fall outside the selected policy type; interpretation must retain that warning."],
        ], [145, 359]),
        H("Why the policy type changes the range"),
        P("The reproducible consumer-demo.md fixture is 434 words and should be submitted as Terms because it spans all eight groups. Its offline Terms baseline also rates 14 of 20 clauses, but its fixed category scope produces 0–42.5 rather than the printed Privacy range of 0–52. Both results are provisional. The screenshot remains valid evidence of what the application displayed; it is not relabeled to imply a different run."),
        H("What a consumer can learn"),
        P("A no-sale promise appears alongside advertiser-sharing permission, so a user can distinguish sale from sharing. Cancellation rights and non-refundable payments are separate findings rather than a single reassuring claim. Perpetual content licensing and AI-training permissions remain prominent concerns, while deletion deadlines and consent-based cookies remain visible protections. The six nuanced clauses are review questions, not invented findings or guarantees about company intent."),
        P("The supplied print does not show accepted AI provenance. It therefore supports a working rule-based report and graceful fallback, not successful live Hybrid acceptance, measured accuracy or complete lifecycle verification."),
    ], True),
    ("8.1 Saved overview and uncertainty", [
        F("overview", "Figure 3. Supplied output, page 1: Hybrid was requested, Rules-only was saved, and 70% coverage is displayed. The Privacy scope warning and provisional 0–52 range are preserved.", 580),
    ], False),
    ("8.2 Concerns and favorable findings", [
        F("highlights", "Figure 4. Supplied output, page 2: perpetual licensing, AI training and changes without notice lead the concern list. Favorable findings are presented separately and do not cancel these concerns.", 580),
    ], False),
    ("8.3 Exact evidence and unresolved wording", [
        F("evidence", "Figure 5. Supplied output, page 4: exact rule matches carry source labels and evidence links. The longer advertiser-transfer clause remains explicitly unresolved instead of receiving an unsupported rating.", 575),
    ], False),
    ("8.4 Navigation, filters and saved actions", [
        F("filters", "Figure 6. Supplied output, page 3: grouped filters and clause cards support focused inspection of the saved result.", 322),
        F("controls", "Figure 7. Supplied output, page 5: paginated clauses and collapsed source/version/comparison/deletion controls keep secondary actions accessible.", 190),
        P("The redesign places the consumer overview before technical detail. Evidence dialogs restore focus on close, severity is conveyed in text as well as color, and narrow-screen layout was checked with Playwright. A printed page cannot demonstrate keyboard behavior; that evidence comes from the recorded browser checks."),
    ], False),
    ("9. Verification and observed evidence", [
        P("The following checks were recorded during implementation. They were not rerun merely to write this report. Test groups can overlap and should not be summed into a claim of independent acceptance coverage. Passing offline checks supports the tested contract, not every deployment or semantic interpretation. [5]"),
        T(["Evidence", "Result and boundary"], [
            ["Contract / ledger / review / persistence / versions", "71 tests passed after the allowance change, including shared totals, retries, stricter account limits and the consumer fixture."],
            ["Groq transport and child lifetime", "29 tests passed with fake HTTP/local child processes; no live provider inference."],
            ["Frontend and browser", "TypeScript/Vite build, 22 Node tests and Playwright smoke passed; API responses were mocked for browser checks."],
            ["Demo setup", "Nine fake-provider tests passed for renewal, identity preservation, failure handling and secret-safe errors."],
            ["Database setup", "User-observed MySQL 8.0.28 logins and both schema checks passed at revision 20261029_lifecycle."],
            ["Runtime print", "Shows 14 rule findings, explicit AI fallback, scope warning and a deletion receipt; not broad acceptance."],
        ], [175, 329]),
        F("deletion", "Figure 8. Supplied output, page 6: a deletion status receipt is visible. This screenshot alone does not establish cleanup correctness under a crash or simultaneous access.", 160),
        P("Frozen plan/source hashes and document-size checks also passed. Outstanding work includes complete real-MySQL worker/restart/two-user/race verification, current live Hybrid evidence, human-labeled evaluation and the fixed performance benchmark."),
    ], True),
    ("10. Limitations and future improvements", [
        H("10.1 Interpretation and coverage"),
        P("The current rules recognize a limited set of supported engineering statements. The reviewed exact-match library is still empty, and no human-labeled held-out evaluation has established precision, recall, macro F1 or semantic coverage. The 70% demo coverage describes one deliberately constructed input; it must not be reported as 70% accuracy or expected performance on arbitrary policies."),
        P("Context handling is currently too broad. A definition or exception anywhere can guard all local rules, while the AI request collects guarded paragraphs across the document. The reported 97-clause policy still produces no AI batch under the modest allowance increase. Fixing this requires section-aware context and relevant definitions without dropping qualifications; raising a budget alone does not resolve the structural problem."),
        H("10.2 AI, deployment and operations"),
        P("AI review is bounded by consent, expiring readiness, provider access, account usage and per-analysis limits. Twelve selected clauses cannot establish high whole-policy coverage for a long document without stronger local rules. Provider refusals, timeouts and validation rejection may preserve a Rules-only result. The supplied output does not establish successful live AI value."),
        P("The application uses a single API process and normal worker for the local demo. API rate counters are not distributed. Public HTTPS deployment would require suitable TLS/cookie configuration, operational monitoring, backup/restore arrangements and broader security testing. Dynamic or blocked websites may require pasted text. English is the supported interpretation scope; jurisdiction-specific enforceability is not determined."),
        H("10.3 Prioritized next work"),
        B("Complete real-stack isolation, restart, race and cleanup acceptance with reproducible records."),
        B("Improve context selection and human-reviewed rules before promising broader policy coverage."),
        B("Evaluate both modes using the frozen plan's labeled data, held-out families and adversarial pairs."),
        B("Measure the prescribed 10,000-word benchmark using warm-up plus 20 runs; report median/p95 and separate network, queue and AI time."),
        B("Retain a small current live Hybrid example with accepted source/model evidence, without unnecessary repeated provider calls."),
    ], True),
    ("11. Conclusion", [
        P("TermsLens demonstrates how a full-stack application can make consumer-policy analysis inspectable. The browser organizes risks, protections and unknowns; the API enforces access and input boundaries; the worker performs durable analysis; and the database retains source versions, evidence and processing state. Optional AI is contained by consent, quotas, schema validation and exact-source checks."),
        P("The supplied output establishes a useful local baseline: fourteen supported findings and clear navigation over a twenty-clause consumer example. It also exposes the limits honestly through Rules-only fallback, unresolved wording and a scope warning. The distinction matters because a completed request, a high coverage percentage and a working provider configuration answer different questions."),
        P("As a mini project, the work demonstrates application architecture, relational persistence, security controls, failure recovery and user-centered presentation. It is not yet an independently validated general policy assessor. The next improvements should strengthen real-stack acceptance and interpretation quality while preserving exact evidence and visible uncertainty."),
        H("Practical demonstration sequence"),
        B("Start the API and one normal worker; keep both terminals open."),
        B("Log in, paste consumer-demo.md and select Terms so all eight groups are in scope."),
        B("Use Hybrid only when available and check the explicit Groq consent box."),
        B("Inspect concerns, protections, unresolved clauses and exact quotes before interpreting score bounds."),
        B("Check requested versus saved mode and any AI source/model labels; reopen the saved report from history."),
        P("The report is based on the supplied cover/output PDFs, the frozen project plan and maintained implementation evidence as of 30 October 2026. It does not mark pending acceptance gates complete."),
    ], True),
    ("Appendix A. Running the demonstration", [
        P("Use two PowerShell terminals at the project root. The existing MySQL setup and migrations are prerequisites; fresh-clone setup is documented in README.md and RUNNING.md. Stop an older API/worker with Ctrl+C before restarting to load code or environment changes. Never paste passwords or keys into these commands."),
        H("Terminal 1 — API"),
        C("$env:TERMSLENS_GROQ_ENABLED = 'true'\n.\\.venv\\Scripts\\python.exe -m uvicorn app.main:app `\n  --app-dir apps/api --host 127.0.0.1 --port 8000"),
        H("Terminal 2 — worker"),
        C("$env:TERMSLENS_GROQ_ENABLED = 'true'\nPush-Location apps/api\n..\\..\\.venv\\Scripts\\python.exe -m app.worker"),
        H("Terminal 3 — health and sample"),
        C("(Invoke-RestMethod 'http://127.0.0.1:8000/api/health').external_review\nGet-Content -LiteralPath './eval/fixtures/consumer-demo.md' `\n  -Raw -Encoding UTF8 | Set-Clipboard"),
        P("Open http://127.0.0.1:8000. Choose Terms, paste the clipboard content, select Hybrid and give consent. The enable variable must be set in each launch terminal because PowerShell environment changes are process-local. An enable flag alone does not replace readiness or available quota."),
        H("If AI is unavailable"),
        P("Read external_review.reason. For disabled, restart the corresponding process after setting its flag. For expired readiness, use the documented explicit renewal only with current account evidence. Do not delete the quota ledger or edit confirmation timestamps to force availability. The worker's waiting message is normal; --test deliberately disables provider settings."),
        P("Full commands, expected findings and troubleshooting remain in README.md, DEMO.md and architecture-debugging.md. Saved reports retain the original result; restarting a process does not retroactively add AI findings."),
    ], True),
    ("References and source notes", [
        P("Project sources were used as the primary evidence. Figure captions identify original supplied output pages; diagrams were authored from code and workflow contracts. No new provider call was made to prepare the report."),
        P("[1] plan-proposal.md. Frozen TermsLens requirements, scoring rubric, operating bounds and five delivery gates."),
        P("[2] implementation-workflow.md. User flow, F1–F8 persistence boundaries, job states, recovery and verification contracts."),
        P("[3] architecture-debugging.md; url-fetching.md; versions.md; comparisons-deletion.md. Maintained architecture, diagnostics and lifecycle documentation."),
        P("[4] Application source: apps/api/app/{main,security,config,analysis,scoring,ai_contract,ai_ledger,ai_review,groq_config,groq_transport,worker}.py; URL modules; apps/web/src report and workflow components."),
        P("[5] verification-checklist.md and transfercontext.md. Recorded checks and pending runtime, provider, evaluation and benchmark acceptance, as inspected on 30 October 2026."),
        P("[6] DEMO.md; README.md; RUNNING.md; eval/fixtures/consumer-demo.md. Demo operation, synthetic example and expected local baseline."),
        P("[7] sampe output.pdf. User-supplied six-page browser print dated 30 October 2026, approximately 09:20. Used for Figures 3–8 without changing displayed results."),
        P("[8] Assignment_10_Report.pdf, page 1. User-supplied cover reference. Student, faculty, course and institutional details retained; heading adapted to Mini Project. The original file remains unchanged."),
        H("Artifact reproducibility"),
        P("scripts/build_project_report.py regenerates this PDF and its editable Word companion using the supplied PDFs and installed local document libraries. Intermediate images and a source-hash manifest are stored under private var/report-build. The report source contains no credentials. The frozen plan, original assignment report and supplied sample output are preserved."),
        C("python scripts/build_project_report.py\n.\\.venv\\Scripts\\python.exe scripts/check_context.py"),
    ], True),
]


def styles():
    pdfmetrics.registerFont(TTFont("ReportBody", str(FONT / "arial.ttf")))
    pdfmetrics.registerFont(TTFont("ReportBold", str(FONT / "arialbd.ttf")))
    return {
        "p": ParagraphStyle("Body", fontName="ReportBody", fontSize=10.3, leading=14.5, textColor=colors.HexColor(INK), spaceAfter=9),
        "h": ParagraphStyle("Subhead", fontName="ReportBold", fontSize=11.5, leading=15, textColor=colors.HexColor(GREEN), spaceBefore=5, spaceAfter=7, keepWithNext=True),
        "title": ParagraphStyle("Chapter", fontName="ReportBold", fontSize=19, leading=23, textColor=colors.HexColor(INK), spaceAfter=17, keepWithNext=True),
        "bullet": ParagraphStyle("Bullet", fontName="ReportBody", fontSize=10.3, leading=14.2, leftIndent=13, firstLineIndent=-10, spaceAfter=6),
        "cell": ParagraphStyle("Cell", fontName="ReportBody", fontSize=9.2, leading=12.2),
        "th": ParagraphStyle("TableHead", fontName="ReportBold", fontSize=9.3, leading=12.5, textColor=colors.white),
        "caption": ParagraphStyle("Caption", fontName="ReportBody", fontSize=9, leading=12.3, textColor=colors.HexColor("#50665d"), spaceBefore=6, spaceAfter=10),
        "code": ParagraphStyle("Code", fontName="Courier", fontSize=8.6, leading=12.4, backColor=colors.HexColor(LIGHT), borderPadding=9, spaceBefore=6, spaceAfter=14),
        "toc": ParagraphStyle("TOC", fontName="ReportBody", fontSize=10.5, leading=18, spaceBefore=7),
    }


class ReportDoc(BaseDocTemplate):
    def __init__(self, filename):
        super().__init__(str(filename), pagesize=(612, 792), leftMargin=54, rightMargin=54, topMargin=52, bottomMargin=46,
                         title="Mini Project — TermsLens", author="KAMAL SUNDAR SK [26120137]", subject="Full Stack Programming mini project report")
        frame = Frame(54, 46, 504, 694, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
        self.addPageTemplates(PageTemplate(id="report", frames=[frame], onPage=self.decorate))
        self.heading_pages = {}

    def decorate(self, canvas, doc):
        if doc.page == 1:
            canvas.drawImage(str(WORK / "cover.png"), 0, 0, width=612, height=792)
            return
        canvas.setStrokeColor(colors.HexColor("#cbd8cf"))
        canvas.line(54, 754, 558, 754)
        canvas.line(54, 36, 558, 36)
        canvas.setFont("ReportBody", 8)
        canvas.setFillColor(colors.HexColor("#50665d"))
        canvas.drawString(54, 764, "MINI PROJECT  |  TERMSLENS")
        canvas.drawRightString(558, 764, "Full Stack Programming")
        canvas.drawString(54, 23, "KAMAL SUNDAR SK  •  26120137")
        canvas.drawRightString(558, 23, str(doc.page))

    def afterFlowable(self, flow):
        if hasattr(flow, "report_heading"):
            title, key = flow.report_heading
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(title, key, 0, False)
            self.notify("TOCEntry", (0, title, self.page, key))
            self.heading_pages[title] = self.page


def build_pdf():
    s = styles()
    story = [Spacer(1, 1), PageBreak()]
    for index, (title, blocks, in_toc) in enumerate(UNITS):
        if index: story.append(PageBreak())
        head = Paragraph(escape(title), s["title"])
        if in_toc: head.report_heading = (title, f"section-{index}")
        story.append(head)
        for block in blocks:
            kind = block[0]
            if kind in {"p", "h", "bullet", "code"}:
                value = escape(block[1])
                if kind == "bullet": value = "• " + value
                if kind == "code": value = value.replace("\n", "<br/>")
                story.append(Paragraph(value, s[kind]))
            elif kind == "table":
                _, headers, rows, widths = block
                data = [[Paragraph(escape(v), s["th"]) for v in headers]]
                data += [[Paragraph(escape(v), s["cell"]) for v in row] for row in rows]
                table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT", spaceAfter=12)
                table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor(GREEN)),
                    ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor(LIGHT)]),
                    ("VALIGN", (0,0), (-1,-1), "TOP"), ("LEFTPADDING", (0,0), (-1,-1), 8),
                    ("RIGHTPADDING", (0,0), (-1,-1), 8), ("TOPPADDING", (0,0), (-1,-1), 7),
                    ("BOTTOMPADDING", (0,0), (-1,-1), 7), ("LINEBELOW", (0,-1), (-1,-1), .5, colors.HexColor("#cbd8cf"))]))
                story.append(table)
            elif kind == "figure":
                _, name, caption, height = block
                path = WORK / f"{name}.png"
                with PILImage.open(path) as im: w, h = im.size
                scale = min(504/w, height/h)
                image = Image(str(path), width=w*scale, height=h*scale)
                image.hAlign = "CENTER"
                story.extend([image, Paragraph(escape(caption), s["caption"])])
            elif kind == "toc":
                toc = TableOfContents()
                toc.levelStyles = [s["toc"]]
                story.append(toc)
                story.append(Spacer(1, 24))
                story.append(Paragraph("Figures 1–2 are architecture/recovery diagrams. Figures 3–8 reproduce the supplied output. Subsections continue within each numbered chapter.", s["caption"]))
    doc = ReportDoc(OUT / "TermsLens_Mini_Project_Report.pdf")
    doc.multiBuild(story)
    return doc.heading_pages


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run._r.addnext(field)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.55)
    section.left_margin = section.right_margin = Inches(.7)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10.5)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size in [("Title", 22), ("Heading 1", 18), ("Heading 2", 12)]:
        style = doc.styles[style_name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor.from_string("203B34")
    # Cover is a faithful image of the supplied first page with only its title area adapted.
    cover = doc.add_paragraph()
    cover.paragraph_format.space_after = Pt(0)
    cover.alignment = 1
    cover.add_run().add_picture(str(WORK / "cover.png"), height=Inches(9.5))
    header = section.header.paragraphs[0]
    header.text = "MINI PROJECT  |  TERMSLENS"
    header.style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.text = "KAMAL SUNDAR SK  •  26120137                                      "
    add_field(footer, "PAGE")
    doc.core_properties.title = "Mini Project — TermsLens"
    doc.core_properties.author = "KAMAL SUNDAR SK [26120137]"
    for title, blocks, in_toc in UNITS:
        doc.add_page_break()
        doc.add_heading(title, level=1 if in_toc else 2)
        for block in blocks:
            kind = block[0]
            if kind == "toc":
                doc.add_paragraph("Chapters (use Word's table-of-contents tools to add page numbers after editing):")
                for heading, _, include in UNITS:
                    if include: doc.add_paragraph(heading)
            elif kind in {"p", "h", "bullet", "code"}:
                if kind == "h": doc.add_heading(block[1], level=2)
                else:
                    p = doc.add_paragraph(block[1], style="List Bullet" if kind == "bullet" else None)
                    if kind == "code":
                        for run in p.runs: run.font.name, run.font.size = "Consolas", Pt(9)
            elif kind == "table":
                _, headers, rows, widths = block
                table = doc.add_table(rows=1, cols=len(headers))
                table.style = "Light Shading Accent 1"
                for i, value in enumerate(headers): table.rows[0].cells[i].text = value
                for row in rows:
                    cells = table.add_row().cells
                    for i, value in enumerate(row): cells[i].text = value
                for row in table.rows:
                    for i, cell in enumerate(row.cells):
                        if widths: cell.width = Inches(widths[i]/72)
                        for p in cell.paragraphs:
                            for run in p.runs: run.font.size = Pt(9.5)
                doc.add_paragraph()
            elif kind == "figure":
                _, name, caption, height = block
                path = WORK / f"{name}.png"
                with PILImage.open(path) as im: w, h = im.size
                scale = min(7/w, height/72/h)
                doc.add_picture(str(path), width=Inches(w*scale))
                doc.add_paragraph(caption, style="Caption")
    doc.save(OUT / "TermsLens_Mini_Project_Report.docx")


def main():
    WORK.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(exist_ok=True)
    before = {name: sha256((ROOT / name).read_bytes()).hexdigest() for name in SOURCE_FILES}
    prepare_images()
    pages = build_pdf()
    build_docx()
    after = {name: sha256((ROOT / name).read_bytes()).hexdigest() for name in SOURCE_FILES}
    assert before == after, "Source input changed"
    manifest = {"sources": before, "section_pages": pages, "inference_calls": 0,
                "outputs": ["reports/TermsLens_Mini_Project_Report.pdf", "reports/TermsLens_Mini_Project_Report.docx"]}
    (WORK / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(json.dumps(manifest["outputs"]))


if __name__ == "__main__": main()
